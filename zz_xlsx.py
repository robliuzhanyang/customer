import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
rows=json.load(open('/workspace/qcs_new_1.json'))+json.load(open('/workspace/qcs_new_2.json'))
chk={x['n']:x['c'] for x in json.load(open('/workspace/qcs_check.json'))}
H=['序号','企业名称','法定代表人','注册资本','成立时间','所在地区','已有资质','推荐理由','联系方式数','是否已入CRM']
F='Arial'
thin=Side(style='thin',color='D9DEE7')
bd=Border(left=thin,right=thin,top=thin,bottom=thin)
hf=PatternFill('solid',fgColor='1F4E79')
z1=PatternFill('solid',fgColor='FFFFFF')
z2=PatternFill('solid',fgColor='F7F9FC')
wb=Workbook();ws=wb.active;ws.title='资质新办-贵州'
for i,h in enumerate(H,1):
    c=ws.cell(1,i,h);c.font=Font(name=F,bold=True,color='FFFFFF',size=10);c.fill=hf;c.border=bd;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
ws.row_dimensions[1].height=26
W=[6,38,11,16,13,30,26,46,10,12]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
for i,r in enumerate(rows,1):
    n=r.get('n','');c=chk.get(n,-1)
    vals=[i,n,r.get('l',''),r.get('cap',''),r.get('f',''),r.get('rg',''),r.get('aids',''),'；'.join(r.get('tags',[])),r.get('tel',''),'是' if c and c>0 else ('否' if c==0 else '未知')]
    fill=z1 if i%2 else z2
    for j,v in enumerate(vals,1):
        cell=ws.cell(i+1,j,v);cell.font=Font(name=F,size=9);cell.fill=fill;cell.border=bd;cell.alignment=Alignment(vertical='center',wrap_text=(j in (2,6,7,8)))
ws.freeze_panes='A2'
ws.auto_filter.ref='A1:J%d'%(len(rows)+1)
wb.save('/workspace/通讯录输出/资质新办_贵州_50条.xlsx')
print('rows',len(rows),'inCRM',sum(1 for r in rows if chk.get(r.get('n',''),0)>0))