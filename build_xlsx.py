import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

rows = []
with open('/workspace/data_page1.json', encoding='utf-8') as f:
    rows += [(1, r) for r in json.load(f)]
with open('/workspace/data_page2.json', encoding='utf-8') as f:
    rows += [(2, r) for r in json.load(f)]

headers = ['序号', '页码', '客户ID', '客户名称', '公司名称', '公司法人', '负责人',
           '手机', 'Hot取号', '电话', '企业地址', '地区', 'Hot联系方式(打码)', '备注']

wb = Workbook()
ws = wb.active
ws.title = '我的客户_前20条'

thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill('solid', fgColor='1F4E79')

for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True, color='FFFFFF', size=11)
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border

for i, (page, r) in enumerate(rows, start=2):
    note = []
    if r.get('hotBtn') == 1:
        note.append('Hot已解锁')
    elif r.get('hotmask') in ('', '--'):
        note.append('无Hot联系方式')
    else:
        note.append('Hot仅显示姓名，无解锁按钮')
    if r.get('mobile') in ('', '--'):
        note.append('平台无手机号')
    if r.get('tel') in ('', '--'):
        note.append('无固定电话')

    vals = [
        i - 1, page, r['key'], r['name'], r['companyD'], r['legal'], r['charger'],
        r['mobile'], r.get('hotNum', ''), r['tel'], r['address'], r['region'],
        r['hotmask'], '；'.join(note)
    ]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.border = border
        cell.alignment = Alignment(vertical='center', wrap_text=(c in (5, 11, 13, 14)))

widths = [6, 6, 26, 34, 34, 12, 10, 15, 15, 16, 62, 34, 20, 34]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'

out = '/workspace/探迹CRM_我的客户_前20条.xlsx'
wb.save(out)

with_mobile = sum(1 for _, r in rows if r.get('mobile') not in ('', '--'))
hot = sum(1 for _, r in rows if r.get('hotNum'))
print('rows', len(rows), 'mobile', with_mobile, 'hot', hot)
print('saved', out)
