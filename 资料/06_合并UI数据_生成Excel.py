# -*- coding: utf-8 -*-
"""
合并 ui_chunks/p01..p10.json（UI 逐条点击采集的前 10 页数据）：
  1) 校验条数与重复客户ID
  2) 输出合并 JSON
  3) 输出汇总 Excel（openpyxl）
用法：python 06_合并UI数据_生成Excel.py
"""
import glob
import json
import os

CHUNK_DIR = r"G:\1\探迹客户数据\工作目录\ui_chunks"
OUT_JSON = r"G:\1\探迹客户数据\结果目录\探迹客户_前10页_明细.json"
OUT_XLSX = r"G:\1\探迹客户数据\结果目录\探迹客户_前10页.xlsx"

HEADERS = ["序号", "页", "客户名称", "公司名称", "公司法人", "负责人",
           "手机", "电话", "企业地址", "地区", "Hot联系方式", "客户ID"]
KEYS = ["no", "page", "name", "company", "legal", "charger",
        "mobile", "tel", "address", "region", "hot", "key"]


def cell(v):
    if v is None or v == "":
        return "--"
    return str(v)


def main():
    files = sorted(glob.glob(os.path.join(CHUNK_DIR, "p*.json")))
    recs = []
    for fp in files:
        with open(fp, "r", encoding="utf-8") as f:
            recs.extend(json.load(f))

    keys = [r.get("key") for r in recs]
    dup = len(keys) - len(set(keys))
    print("分片文件: %d 个, 记录: %d 条, 重复客户ID: %d" % (len(files), len(recs), dup))

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(recs, f, ensure_ascii=False, indent=1)

    with_mobile = sum(1 for r in recs if r.get("mobile") and r["mobile"] != "--")
    with_addr = sum(1 for r in recs if r.get("address"))
    print("有手机号: %d 条, 有地址: %d 条" % (with_mobile, with_addr))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("未安装 openpyxl，跳过 Excel")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "客户数据_前10页"
    ws.append(HEADERS)
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="4472C4")
    for c in ws[1]:
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")
    for r in recs:
        ws.append([cell(r.get(k)) for k in KEYS])
    widths = [6, 5, 34, 34, 12, 10, 15, 16, 46, 26, 22, 26]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)
    print("Excel 已生成: %s" % OUT_XLSX)


if __name__ == "__main__":
    main()
