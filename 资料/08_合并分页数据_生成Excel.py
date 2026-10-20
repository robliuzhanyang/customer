# -*- coding: utf-8 -*-
"""
合并 ui_pages/pg*.json（逐页 UI 采集分片）→ 汇总 JSON + Excel（可随时重跑，做阶段存档）。

列：序号/页/客户名称/公司名称/公司法人/负责人/手机/Hot取号/电话/企业地址/地区/Hot联系方式/客户ID/备注
用法：python 08_合并分页数据_生成Excel.py [输出Excel路径]
"""
import glob
import json
import os
import sys

CHUNK_DIR = r"G:\1\探迹客户数据\工作目录\ui_pages"
OUT_JSON = r"G:\1\探迹客户数据\结果目录\探迹客户_采集总表.json"
OUT_XLSX = sys.argv[1] if len(sys.argv) > 1 else r"G:\1\探迹客户数据\结果目录\探迹客户_采集总表.xlsx"

HEADERS = ["序号", "页", "客户名称", "公司名称", "公司法人", "负责人", "手机", "Hot取号",
           "电话", "企业地址", "地区", "Hot联系方式", "客户ID", "备注"]
KEYS = ["no", "page", "name", "company", "legal", "charger", "mobile", "hotRevealed",
        "tel", "address", "region", "hot", "key"]


def has_num(v):
    return bool(v) and v != "--" and any(c.isdigit() for c in str(v))


def main():
    files = sorted(glob.glob(os.path.join(CHUNK_DIR, "pg*.json")))
    recs = []
    for fp in files:
        with open(fp, "r", encoding="utf-8") as f:
            r = json.load(f)
            if isinstance(r, list):
                recs.extend(r)
    recs.sort(key=lambda x: x.get("no", 0))
    keys = [r.get("key") for r in recs]
    dup = len(keys) - len(set(keys))

    for r in recs:
        rm = ""
        if not has_num(r.get("mobile")):
            if has_num(r.get("hotRevealed")):
                rm = "手机缺失，已用列表Hot取号"
            elif has_num(r.get("tel")):
                rm = "平台无手机号，仅有座机"
            elif r.get("hot") and r["hot"] != "--":
                rm = "平台仅有Hot姓名，无号码"
            else:
                rm = "平台无任何号码"
        r["remark"] = rm

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(recs, f, ensure_ascii=False, indent=1)

    print("分片 %d 个, 记录 %d 条, 重复ID %d" % (len(files), len(recs), dup))
    print("有手机 %d, 有Hot取号 %d, 有地址 %d" % (
        sum(1 for r in recs if has_num(r.get("mobile"))),
        sum(1 for r in recs if has_num(r.get("hotRevealed"))),
        sum(1 for r in recs if r.get("address") and r["address"] != "--")))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("未安装 openpyxl")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "客户数据"
    ws.append(HEADERS)
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="4472C4")
    for c in ws[1]:
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")
    for r in recs:
        ws.append([(r.get(k) if r.get(k) not in (None, "") else "--") for k in KEYS] + [r.get("remark", "")])
    widths = [6, 5, 34, 34, 12, 10, 15, 16, 16, 46, 26, 22, 26, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)
    print("Excel -> %s" % OUT_XLSX)


if __name__ == "__main__":
    main()
