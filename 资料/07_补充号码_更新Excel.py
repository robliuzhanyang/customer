# -*- coding: utf-8 -*-
"""
在「前10页明细」基础上应用 Hot 联系方式补录结果，重新生成 Excel（含备注列）。

补录来源：ui_chunks/hot_supplement.json（点击列表 Hot联系方式 按钮获取的真实号码）
用法：python 07_补充号码_更新Excel.py
"""
import json
import os

SRC_JSON = r"G:\1\探迹客户数据\结果目录\探迹客户_前10页_明细.json"
SUP_JSON = r"G:\1\探迹客户数据\工作目录\ui_chunks\hot_supplement.json"
OUT_XLSX = r"G:\1\探迹客户数据\结果目录\探迹客户_前10页.xlsx"
OUT_JSON = r"G:\1\探迹客户数据\结果目录\探迹客户_前10页_明细.json"

HEADERS = ["序号", "页", "客户名称", "公司名称", "公司法人", "负责人", "手机", "电话",
           "企业地址", "地区", "Hot联系方式", "客户ID", "备注"]
KEYS = ["no", "page", "name", "company", "legal", "charger", "mobile", "tel",
        "address", "region", "hot", "key"]


def has_number(v):
    return bool(v) and v != "--" and any(ch.isdigit() for ch in str(v))


def main():
    with open(SRC_JSON, "r", encoding="utf-8") as f:
        recs = json.load(f)

    sup = {}
    if os.path.exists(SUP_JSON):
        with open(SUP_JSON, "r", encoding="utf-8") as f:
            for item in json.load(f):
                if isinstance(item, dict) and item.get("numbers"):
                    sup[item["no"]] = item["numbers"][0]

    for r in recs:
        r["remark"] = ""
        no = r.get("no")
        if no in sup:
            num = sup[no]
            name_part = (r.get("hot") or "").split("：")[0].rstrip("：:") or "已获取"
            r["hot"] = "%s：%s" % (name_part, num)
            r["remark"] = "手机号平台未提供，已点击列表Hot联系方式取得号码"
        elif not has_number(r.get("mobile")):
            if has_number(r.get("tel")):
                r["remark"] = "平台无手机号，仅有座机"
            else:
                r["remark"] = "平台无任何号码（电话/手机/Hot 均为空）"

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(recs, f, ensure_ascii=False, indent=1)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("未安装 openpyxl")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "客户数据_前10页"
    ws.append(HEADERS)
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="4472C4")
    for c in ws[1]:
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")
    for r in recs:
        row = []
        for k in KEYS:
            v = r.get(k)
            row.append(v if v not in (None, "") else "--")
        row.append(r.get("remark", ""))
        ws.append(row)
    widths = [6, 5, 34, 34, 12, 10, 15, 16, 46, 26, 24, 26, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)

    filled = sum(1 for r in recs if has_number(r.get("mobile")))
    sup_n = len(sup)
    print("共 %d 条；手机号有效 %d 条；本次补录 %d 条 -> %s" % (len(recs), filled, sup_n, OUT_XLSX))
    for r in recs:
        if r.get("remark"):
            print("  备注 no.%s %s -> %s" % (r["no"], r["name"], r["remark"]))


if __name__ == "__main__":
    main()
