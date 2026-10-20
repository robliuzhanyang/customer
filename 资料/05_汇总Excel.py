# -*- coding: utf-8 -*-
"""
把全量客户 JSON 汇总成一个 Excel（openpyxl）；没装 openpyxl 时自动降级输出 CSV(utf-8-sig)。
用法：python 05_汇总Excel.py [源JSON路径]
"""
import csv
import json
import sys

SRC = sys.argv[1] if len(sys.argv) > 1 else r"G:\1\探迹客户数据\结果目录\客户数据_全量.json"
OUT_XLSX = r"G:\1\探迹客户数据\结果目录\客户数据汇总.xlsx"
OUT_CSV = r"G:\1\探迹客户数据\结果目录\客户数据汇总.csv"

HEADERS = ["序号", "客户名称", "公司名称", "公司法人", "Hot联系方式", "负责人"]
KEYS = ("no", "name", "company", "legal", "hot", "charger")


def main():
    with open(SRC, "r", encoding="utf-8") as f:
        data = json.load(f)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(HEADERS)
            for it in data:
                w.writerow([it.get(k, "") for k in KEYS])
        print("未安装 openpyxl，已降级输出 CSV：%s（%d 条）" % (OUT_CSV, len(data)))
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "客户数据"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    for it in data:
        ws.append([it.get(k, "") for k in KEYS])
    widths = [8, 42, 42, 16, 26, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    wb.save(OUT_XLSX)
    print("完成：Excel 已生成 %s（共 %d 条）" % (OUT_XLSX, len(data)))


if __name__ == "__main__":
    main()
