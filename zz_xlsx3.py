import json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = 'Arial'
thin = Side(style='thin', color='D9DEE7')
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
hf = PatternFill('solid', fgColor='1F4E79')
z1 = PatternFill('solid', fgColor='FFFFFF')
z2 = PatternFill('solid', fgColor='F7F9FC')


def load(paths):
    out, seen = [], set()
    for p in paths:
        if not os.path.exists(p):
            continue
        for r in json.load(open(p)):
            n = (r.get('n') or '').strip()
            if not n or n in seen:
                continue
            seen.add(n)
            out.append(r)
    return out


def build(rows, title, rq_label, out_path):
    H = ['序号', '企业名称', '法定代表人', '注册资本', '成立时间', '所在地区', rq_label, '联系方式数']
    W = [6, 40, 11, 16, 13, 30, 62, 10]
    wb = Workbook(); ws = wb.active; ws.title = title
    for i, h in enumerate(H, 1):
        c = ws.cell(1, i, h)
        c.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        c.fill = hf; c.border = bd
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 26
    for i, w in enumerate(W, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, r in enumerate(rows, 1):
        vals = [i, r.get('n', ''), r.get('l', ''), r.get('cap', ''), r.get('f', ''),
                r.get('rg', ''), (r.get('rq', '') or '').replace(chr(10), '；'), r.get('tel', '')]
        fill = z1 if i % 2 else z2
        for j, v in enumerate(vals, 1):
            c = ws.cell(i + 1, j, v)
            c.font = Font(name=F, size=9); c.fill = fill; c.border = bd
            c.alignment = Alignment(vertical='center', wrap_text=(j in (2, 6, 7)))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:H%d' % (len(rows) + 1)
    wb.save(out_path)
    return len(rows)


d = load(['/workspace/qcs_delay_1.json', '/workspace/qcs_delay_2.json'] +
         ['/workspace/qcs_delay_p%d.json' % i for i in range(2, 7)])
n = load(['/workspace/qcs_new_1.json', '/workspace/qcs_new_2.json'] +
         ['/workspace/qcs_new_p%d.json' % i for i in range(2, 8)])

os.makedirs('/workspace/通讯录输出', exist_ok=True)
a = build(d, '资质延期-贵州', '待延期资质', '/workspace/通讯录输出/资质延期_贵州_%d条.xlsx' % len(d))
b = build(n, '资质新办-贵州', '新办资质/推荐理由', '/workspace/通讯录输出/资质新办_贵州_%d条.xlsx' % len(n))
print('delay', a, 'new', b)
