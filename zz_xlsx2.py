import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
rows=json.load(open('/workspace/qcs_delay_1.json'))+json.load(open('/workspace/qcs_delay_2.json'))
H=['序号','企业名称','法定代表人','注册资本','成立时间','所在地区','待延期资质','联系方式数']
F='Arial'
thin=Side(style='thin',color='D9DEE7')
bd=Border(left=thin,right=thin,top=thin,bottom=thin)
hf=PatternFill('solid',fgColor='1F4E79')
z1=PatternFill('solid',fgColor='FFFFFF')
z2=PatternFill('solid',fgColor='F7F9FC')
wb=Workbook();ws=wb.active;ws.title='资质延期-贵州'
for i,h in enumerate(H,1):
    c=ws.cell(1,i,h);c.font=Font(name=F,bold=True,color='FFFFFF',size=10);c.fill=hf;c.border=bd;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
ws.row_dimensions[1].height=26
W=[6,40,11,16,13,30,60,10]
for i,w in enumerate(W,1): ws.column_dimensions[get_column_letter(i)].width=w
for i,r in enumerate(rows,1):
    vals=[i,r.get('n',''),r.get('l',''),r.get('cap',''),r.get('f',''),r.get('rg',''),(r.get('rq','') or '').replace(chr(10),'；'),r.get('tel','')]
    fill=z1 if i%2 else z2
    for j,v in enumerate(vals,1):
        c=ws.cell(i+1,j,v);c.font=Font(name=F,size=9);c.fill=fill;c.border=bd;c.alignment=Alignment(vertical='center',wrap_text=(j in (2,6,7)))
ws.freeze_panes='A2';ws.auto_filter.ref='A1:H%d'%(len(rows)+1)
wb.save('/workspace/通讯录输出/资质延期_贵州_50条.xlsx')
print('rows',len(rows))