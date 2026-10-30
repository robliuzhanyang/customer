import json, glob
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = 'Arial'
thin = Side(style='thin', color='D9DEE7')
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
hf = PatternFill('solid', fgColor='1F4E79')
z1 = PatternFill('solid', fgColor='FFFFFF')
z2 = PatternFill('solid', fgColor='F7F9FC')


def load(prefix):
    rows, seen = [], set()
    for f in sorted(glob.glob('/workspace/%s_p*.json' % prefix),
                    key=lambda x: int(x.split('_p')[1].split('.')[0])):
        for r in json.load(open(f)):
            n = (r.get('n') or '').strip()
            if not n or n in seen:
                continue
            seen.add(n)
            rows.append(r)
    return rows


def build(rows, sheet, label, out):
    H = ['序号', '企业名称', '法定代表人', '注册资本', '成立时间', '所在地区', label, '联系方式数']
    W = [6, 42, 11, 16, 13, 26, 62, 10]
    wb = Workbook(); ws = wb.active; ws.title = sheet
    for i, h in enumerate(H, 1):
        c = ws.cell(1, i, h)
        c.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        c.fill = hf; c.border = bd
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 26
    for i, w in enumerate(W, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, r in enumerate(rows, 1):
        vals = [i, r.get('n', ''), r.get('l', ''), r.get('cap', ''), r.get('f', ''),
                r.get('rg', ''), (r.get('rq', '') or '').replace(chr(10), '；'), r.get('tel', '')]
        fill = z1 if i % 2 else z2
        for j, v in enumerate(vals, 1):
            c = ws.cell(i + 1, j, v)
            c.font = Font(name=F, size=9); c.fill = fill; c.border = bd
            c.alignment = Alignment(vertical='center', wrap_text=(j in (2, 6, 7)))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:H%d' % (len(rows) + 1)
    wb.save(out)
    return len(rows)


out_dir = '/workspace/通讯录输出/'
specs = [
    ('qcs_upg', '资质升级-贵州', '资质升级/推荐理由', '资质升级_贵州_%d条.xlsx'),
    ('qcs_dyn', '动态核查-贵州', '核查信息', '动态核查_贵州_%d条.xlsx'),
    ('qcs_abn', '资质异常-贵州', '异常信息', '资质异常_贵州_%d条.xlsx'),
    ('qcs_rej', '申请未过', '未过原因/推荐理由', '申请未过_%d条.xlsx'),
]
for prefix, sheet, label, tpl in specs:
    rows = load(prefix)
    p = out_dir + (tpl % len(rows))
    build(rows, sheet, label, p)
    print(prefix, len(rows), p)
